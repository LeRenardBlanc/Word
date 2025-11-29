import os
from fpdf import FPDF
from docx import Document
import markdown
import tempfile
from .models import Note

def export_to_txt(note: Note) -> str:
    """Exports note to a temporary TXT file and returns the path."""
    with tempfile.NamedTemporaryFile(mode='w', delete=False, suffix=".txt", encoding='utf-8') as tmp:
        tmp.write(f"Title: {note.title}\n\n")
        tmp.write(note.content or "")
        return tmp.name

def export_to_md(note: Note) -> str:
    """Exports note to a temporary Markdown file and returns the path."""
    with tempfile.NamedTemporaryFile(mode='w', delete=False, suffix=".md", encoding='utf-8') as tmp:
        tmp.write(f"# {note.title}\n\n")
        tmp.write(note.content or "")
        return tmp.name

def export_to_docx(note: Note) -> str:
    """Exports note to a temporary DOCX file and returns the path."""
    doc = Document()
    doc.add_heading(note.title, 0)
    doc.add_paragraph(note.content or "")
    
    fd, path = tempfile.mkstemp(suffix=".docx")
    os.close(fd)
    doc.save(path)
    return path

def export_to_pdf(note: Note) -> str:
    """Exports note to a temporary PDF file and returns the path."""
    class PDF(FPDF):
        def header(self):
            self.set_font('Arial', 'B', 12)
            self.cell(0, 10, 'Note Export', 0, 1, 'C')

    pdf = PDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)
    
    # Add title
    pdf.set_font("Arial", 'B', 16)
    pdf.multi_cell(0, 10, note.title.encode('latin-1', 'replace').decode('latin-1'))
    pdf.ln(5)
    
    # Add content
    pdf.set_font("Arial", size=12)
    # FPDF doesn't handle unicode well by default without a font file, so we do basic replacement
    # In a real app, we'd load a unicode font like DejaVuSans
    safe_content = (note.content or "").encode('latin-1', 'replace').decode('latin-1')
    pdf.multi_cell(0, 10, safe_content)
    
    fd, path = tempfile.mkstemp(suffix=".pdf")
    os.close(fd)
    pdf.output(path)
    return path
